from __future__ import annotations

import csv
import json
import logging
from pathlib import Path
from typing import Any

from app.models.scrape_result import ScrapeResult, ExportFormat

logger = logging.getLogger("scraper.export")


class Exporter:
    def __init__(self, output_dir: str = "output"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export(
        self,
        result: ScrapeResult,
        name: str = "scrape_result",
        fmt: ExportFormat = ExportFormat.JSON,
    ) -> str:
        exporters = {
            ExportFormat.JSON: self._to_json,
            ExportFormat.CSV: self._to_csv,
            ExportFormat.EXCEL: self._to_excel,
        }
        handler = exporters.get(fmt)
        if handler is None:
            raise ValueError(f"Formato non supportato: {fmt}")
        return handler(result, name)

    def export_all(
        self, result: ScrapeResult, name: str = "scrape_result"
    ) -> dict[str, str]:
        paths = {}
        for fmt in ExportFormat:
            try:
                paths[fmt.value] = self.export(result, name, fmt)
            except Exception as e:
                logger.warning("Export %s fallito: %s", fmt.value, e)
        return paths

    def _to_json(self, result: ScrapeResult, name: str) -> str:
        path = self.output_dir / f"{name}.json"
        data = result.to_dict()
        path.write_text(
            json.dumps(data, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        logger.info("Esportato JSON: %s", path)
        return str(path)

    def _to_csv(self, result: ScrapeResult, name: str) -> str:
        """CSV leggibile: ogni tabella è espansa in righe reali (apribile in Excel)."""
        path = self.output_dir / f"{name}.csv"

        tables = [i for i in result.items if i.type == "table"]
        texts = [i for i in result.items if i.type == "text"]
        links = [i for i in result.items if i.type == "link"]
        images = [i for i in result.items if i.type == "image"]

        with path.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(["WebScraper Pro — export"])
            writer.writerow(["URL", result.url])
            if result.title:
                writer.writerow(["Titolo", result.title])
            writer.writerow([])

            for idx, item in enumerate(tables, 1):
                data = item.content if isinstance(item.content, dict) else {}
                headers = data.get("headers", []) or []
                rows = data.get("rows", []) or []
                a = item.attributes or {}
                parts = [a.get(k) for k in ("league", "stagione", "metrica", "pagina", "kind") if a.get(k)]
                descr = " — ".join(str(p) for p in parts) if parts else (a.get("source") or "")
                label = f"TABELLA {idx}" + (f" — {descr}" if descr else "")
                writer.writerow([label, f"{len(rows)} righe"])
                if headers:
                    writer.writerow(headers)
                for r in rows:
                    writer.writerow(r)
                writer.writerow([])

            if texts:
                writer.writerow(["TESTI"])
                for it in texts:
                    writer.writerow([str(it.content)[:32000]])
                writer.writerow([])

            if links:
                writer.writerow(["LINK (testo)", "URL"])
                for it in links:
                    c = it.content if isinstance(it.content, dict) else {}
                    writer.writerow([c.get("text", ""), c.get("url", "")])
                writer.writerow([])

            if images:
                writer.writerow(["IMMAGINE (alt)", "URL"])
                for it in images:
                    c = it.content if isinstance(it.content, dict) else {}
                    writer.writerow([c.get("alt", ""), c.get("url", "")])

        logger.info("Esportato CSV: %s", path)
        return str(path)

    def _to_excel(self, result: ScrapeResult, name: str) -> str:
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        path = self.output_dir / f"{name}.xlsx"
        wb = openpyxl.Workbook()

        # Foglio riepilogo
        ws = wb.active
        ws.title = "Summary"
        ws.cell(1, 1, "URL").value = result.url
        ws.cell(2, 1, "Title").value = result.title or ""
        ws.cell(3, 1, "Description").value = result.description or ""
        ws.cell(5, 1, "Type").value = "Type"
        ws.cell(5, 2, "Count").value = "Count"
        type_counts: dict[str, int] = {}
        for item in result.items:
            type_counts[item.type] = type_counts.get(item.type, 0) + 1
        for i, (t, c) in enumerate(sorted(type_counts.items()), start=6):
            ws.cell(i, 1, t)
            ws.cell(i, 2, c)

        # Foglio dati
        ws2 = wb.create_sheet("Data")
        ws2.cell(1, 1, "Type")
        ws2.cell(1, 2, "Content")
        ws2.cell(1, 3, "Selector")
        ws2.cell(1, 4, "Attributes")
        header_font = Font(bold=True)
        for col in range(1, 5):
            ws2.cell(1, col).font = header_font

        for i, item in enumerate(result.items, start=2):
            ws2.cell(i, 1, item.type)
            content_str = str(item.content)[:32000]
            ws2.cell(i, 2, content_str)
            ws2.cell(i, 3, item.selector or "")
            ws2.cell(i, 4, str(item.attributes))

        # Foglio tabelle
        tables = [i for i in result.items if i.type == "table"]
        if tables:
            ws3 = wb.create_sheet("Tables")
            row_num = 1
            for ti, table_item in enumerate(tables):
                data = table_item.content
                headers = data.get("headers", [])
                rows = data.get("rows", [])

                ws3.cell(row_num, 1, f"Table {ti + 1}")
                ws3.cell(row_num, 1).font = Font(bold=True, size=12)
                row_num += 1

                if headers:
                    for ci, h in enumerate(headers, start=1):
                        ws3.cell(row_num, ci, h)
                        ws3.cell(row_num, ci).font = header_font
                    row_num += 1

                for r in rows:
                    for ci, val in enumerate(r, start=1):
                        ws3.cell(row_num, ci, val)
                    row_num += 1
                row_num += 2

        wb.save(path)
        logger.info("Esportato Excel: %s", path)
        return str(path)
